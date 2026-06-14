"""Page-by-page document loaders for PDF, DOCX, TXT and Markdown.

Each loader returns a ``ParsedDocument`` whose ``pages`` attribute is a
1-indexed mapping of page number to extracted text::

    {1: "page 1 text", 2: "page 2 text", ...}

PDFs are parsed with PyMuPDF (fitz) when available, falling back to
pdfplumber. Scanned PDFs (pages with no extractable text) can optionally be
run through Tesseract OCR.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}

# Roughly how many characters make up a "page" for paginated-less formats
# (DOCX without page breaks, large TXT/MD files).
CHARS_PER_SYNTHETIC_PAGE = 3000


@dataclass
class ParsedDocument:
    """Internal representation of a parsed document."""

    pages: dict[int, str] = field(default_factory=dict)
    source_path: str | None = None
    page_count: int = 0
    used_ocr: bool = False

    def non_empty_pages(self) -> dict[int, str]:
        return {n: t for n, t in self.pages.items() if t and t.strip()}


class UnsupportedFileTypeError(ValueError):
    pass


def load_document(file_path: str | Path, *, enable_ocr: bool = False) -> ParsedDocument:
    """Load a document from disk into a ``ParsedDocument``."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        doc = _load_pdf(path, enable_ocr=enable_ocr)
    elif ext == ".docx":
        doc = _load_docx(path)
    elif ext in {".txt", ".md", ".markdown"}:
        doc = _load_text(path)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    doc.source_path = str(path)
    doc.page_count = len(doc.pages)
    logger.info("Parsed %s into %d pages (ocr=%s)", path.name, doc.page_count, doc.used_ocr)
    return doc


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _load_pdf(path: Path, *, enable_ocr: bool) -> ParsedDocument:
    try:
        return _load_pdf_pymupdf(path, enable_ocr=enable_ocr)
    except ImportError:
        logger.warning("PyMuPDF unavailable, falling back to pdfplumber")
        return _load_pdf_pdfplumber(path)


def _load_pdf_pymupdf(path: Path, *, enable_ocr: bool) -> ParsedDocument:
    import fitz  # PyMuPDF

    doc = ParsedDocument()
    with fitz.open(path) as pdf:
        for i, page in enumerate(pdf, start=1):
            text = page.get_text("text") or ""
            if not text.strip() and enable_ocr:
                text = _ocr_pdf_page(page)
                if text.strip():
                    doc.used_ocr = True
            doc.pages[i] = text.strip()
    return doc


def _ocr_pdf_page(page) -> str:
    """Render a PDF page to an image and OCR it with Tesseract."""
    try:
        import io

        import pytesseract
        from PIL import Image

        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(img) or ""
    except Exception as exc:  # pragma: no cover - OCR is best-effort
        logger.warning("OCR failed for a page: %s", exc)
        return ""


def _load_pdf_pdfplumber(path: Path) -> ParsedDocument:
    import pdfplumber

    doc = ParsedDocument()
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            doc.pages[i] = (page.extract_text() or "").strip()
    return doc


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def _load_docx(path: Path) -> ParsedDocument:
    import docx  # python-docx

    document = docx.Document(str(path))

    # DOCX has no reliable page concept; split on explicit page breaks when
    # present, otherwise fall back to synthetic pagination by length.
    paragraphs = [p.text for p in document.paragraphs]
    full_text = "\n".join(paragraphs).strip()

    if _has_page_breaks(document):
        return _paginate_on_breaks(document)
    return _paginate_by_length(full_text)


def _has_page_breaks(document) -> bool:
    xml = document.element.xml
    return 'w:type="page"' in xml or "<w:br" in xml and 'type="page"' in xml


def _paginate_on_breaks(document) -> ParsedDocument:
    doc = ParsedDocument()
    page_no = 1
    buffer: list[str] = []
    for para in document.paragraphs:
        if 'w:type="page"' in para._p.xml:
            doc.pages[page_no] = "\n".join(buffer).strip()
            page_no += 1
            buffer = []
        buffer.append(para.text)
    doc.pages[page_no] = "\n".join(buffer).strip()
    return doc


# --------------------------------------------------------------------------- #
# TXT / Markdown
# --------------------------------------------------------------------------- #
def _load_text(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    # Respect form-feed page separators if present.
    if "\f" in text:
        doc = ParsedDocument()
        for i, chunk in enumerate(text.split("\f"), start=1):
            doc.pages[i] = chunk.strip()
        return doc
    return _paginate_by_length(text)


def _paginate_by_length(text: str) -> ParsedDocument:
    """Split a continuous string into synthetic pages on paragraph boundaries."""
    doc = ParsedDocument()
    text = text.strip()
    if not text:
        doc.pages[1] = ""
        return doc

    paragraphs = text.split("\n\n")
    page_no = 1
    buffer = ""
    for para in paragraphs:
        if buffer and len(buffer) + len(para) > CHARS_PER_SYNTHETIC_PAGE:
            doc.pages[page_no] = buffer.strip()
            page_no += 1
            buffer = ""
        buffer += para + "\n\n"
    if buffer.strip():
        doc.pages[page_no] = buffer.strip()
    return doc
